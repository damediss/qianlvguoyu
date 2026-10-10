from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path("/Users/macmain/Documents/qianlvguoyu")
SRC = BASE / "国昱公司核心团队招聘与培养计划.md"
OUT = BASE / "国昱公司核心团队招聘与培养计划.docx"

BLUE = (31, 78, 121)
DEEP_BLUE_HEX = "1F4E79"
LIGHT_BLUE_HEX = "EAF2F8"
LIGHT_GRAY_HEX = "F7F9FC"
TEXT = (31, 31, 31)
GRAY = (89, 89, 89)


def set_cn_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        child = tc_mar.find(qn(f"w:{name}"))
        if child is None:
            child = OxmlElement(f"w:{name}")
            tc_mar.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")


def set_cell_width(cell, cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def dont_split_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("国昱公司核心团队招聘与培养计划")
    set_cn_font(run, 8.5, False, GRAY)


def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.add_run("  |  第 ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = footer.add_run()
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_end)
    footer.add_run(" 页")


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.style = f"Heading {min(level, 3)}"
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 2 else 6)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    size = 15 if level == 2 else 12.5
    set_cn_font(r, size, True, BLUE)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_cn_font(r, 10.5, False, TEXT)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_cn_font(r, 10.2, False, TEXT)
    return p


def add_table(doc, headers, rows):
    widths = table_widths(len(headers), headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row[: len(headers)]):
            cells[i].text = value

    for ri, row in enumerate(table.rows):
        dont_split_row(row)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ci < len(widths):
                set_cell_width(cell, widths[ci])
            if ri == 0:
                set_shading(cell, DEEP_BLUE_HEX)
            elif ri % 2 == 0:
                set_shading(cell, LIGHT_GRAY_HEX)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.06
                if ri == 0 or headers[ci] in {"人数", "对象"}:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_cn_font(
                        run,
                        8.6 if len(headers) >= 4 else 9.0,
                        ri == 0,
                        (255, 255, 255) if ri == 0 else TEXT,
                    )
    repeat_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def table_widths(n, headers):
    if n == 4 and headers[:2] == ["岗位", "人数"]:
        return [3.4, 1.2, 6.0, 6.2]
    if n == 3:
        return [5.2, 5.6, 5.8]
    if n == 2:
        return [4.8, 11.6]
    if n == 4:
        return [3.2, 3.2, 5.0, 5.2]
    return [16.4 / n] * n


def parse_table(lines, start):
    headers = [c.strip() for c in lines[start].strip().strip("|").split("|")]
    rows = []
    i = start + 2
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    return headers, rows, i


def build():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    add_footer(section)
    add_page_number(section)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = lines[0].lstrip("#").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_cn_font(r, 23, True, BLUE)

    i = 1
    meta_lines = []
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("## "):
            break
        if line:
            meta_lines.append(line.rstrip("  "))
        i += 1
    if meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run("｜".join(meta_lines))
        set_cn_font(r, 10, False, GRAY)

    # Opening callout from the first section's key judgment.
    intro = (
        "第一阶段重点不是扩大编制，而是搭建能接住业务需求、快速试错、快速迭代、"
        "并能借助 AI 把项目经验转化为可复用能力的小型核心团队。"
    )
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout.cell(0, 0)
    set_shading(cell, LIGHT_BLUE_HEX)
    set_cell_margins(cell, 140, 170, 140, 170)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(intro)
    set_cn_font(r, 10.8, True, BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*[-: ]+\|", lines[i + 1].strip()):
            headers, rows, i = parse_table(lines, i)
            add_table(doc, headers, rows)
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
        elif line.startswith("- "):
            add_bullet(doc, line[2:].strip())
        else:
            add_para(doc, line)
        i += 1

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
