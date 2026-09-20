# -*- coding: utf-8 -*-
"""Generate training-dept sign-off Word for default-template export."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

out_dir = Path(r"d:\phpApp\qianlvguoyu_saas\docs")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "开班资料默认模板导出_使用说明与测试确认单.docx"

doc = Document()
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def set_run_font(run, name="宋体", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), name)


def add_para(
    text="",
    *,
    size=12,
    bold=False,
    align=None,
    space_before=0,
    space_after=6,
    first_line=0,
    name="宋体",
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.25
    if first_line:
        pf.first_line_indent = Cm(first_line)
    if text:
        run = p.add_run(text)
        set_run_font(run, name=name, size=size, bold=bold)
    return p


def add_heading_cn(text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    return add_para(
        text,
        size=sizes.get(level, 12),
        bold=True,
        space_before=12 if level > 1 else 6,
        space_after=8,
        name="黑体",
    )


def set_cell_text(cell, text, *, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
        tc_pr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "666666")
        borders.append(el)
    tbl_pr.append(borders)


# Cover
add_para("千铝国昱云平台", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="黑体", space_after=4)
add_para(
    "开班资料「默认模板导出」功能",
    size=18,
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    name="黑体",
    space_before=8,
    space_after=4,
)
add_para("使用说明与测试确认单", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, name="黑体", space_after=12)
add_para(
    f"文档编号：TRN-EXP-20260915　　版本：V1.0　　日期：{date.today().isoformat()}",
    size=10.5,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_after=4,
)
add_para("适用范围：培训部门 · 开班准备 · 班级资料收集", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

meta = doc.add_table(rows=4, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(meta)
meta_data = [
    ("编制部门", "技术研发", "确认部门", "培训部门"),
    ("编制人", "", "确认人", ""),
    ("编制日期", date.today().isoformat(), "确认日期", ""),
    ("关联问题单", "《文档导出问题.docx》", "结论", "☐通过　☐有条件通过　☐不通过"),
]
for i, row in enumerate(meta_data):
    for j, val in enumerate(row):
        set_cell_text(
            meta.cell(i, j),
            val,
            bold=(j % 2 == 0),
            size=10.5,
            fill="F5F5F5" if j % 2 == 0 else None,
        )

add_para("", space_after=8)

# 1
add_heading_cn("一、背景与目的", 1)
add_para(
    "培训部门在《文档导出问题.docx》中反馈了「默认模板导出」三类材料问题"
    "（工作年限证明页边距与排版、职业技能评价登记表工种/申报级别、职业技能等级认定申请表字段与等级）。"
    "上述问题已完成修复并上线。本文件用于说明正确使用方法，并提供验收测试项，供培训部门签字确认。",
    first_line=0.74,
    size=12,
    space_after=8,
)

# 2
add_heading_cn("二、问题闭环清单（均已解决）", 1)
t1 = doc.add_table(rows=4, cols=4)
set_table_borders(t1)
for j, h in enumerate(["序号", "原问题描述", "修复要点", "状态"]):
    set_cell_text(t1.cell(0, j), h, bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9E2F3")

rows = [
    (
        "1",
        "工作证明（工作年限证明）纸张页边距不符；导出排版不正确。要求：上3.7cm、下3.5cm、左2.8cm、右2.6cm。",
        "更新模板《工作年限证明.docx》页边距为 3.7/3.5/2.8/2.6 cm；优化正文排版；证明人/电话优先取班级单位联系人。",
        "已解决",
    ),
    (
        "2",
        "职业技能评价登记表中工种不能一起导出需手工改；申报级别初始化多为「五级/初级工」需手工改。",
        "未单独预约时回退班级工种与班级等级；纠正库默认 req_level=0 误写五级；按是否有上一等级证书勾选晋级/初次。",
        "已解决",
    ),
    (
        "3",
        "认定申请表：毕业时间、电话、地址、邮编未导出；原等级有问题；简历导出错误；认定等级错误。",
        "补齐毕业时间与企业电话/地址/邮编回退；现工种等级无证据时留空；工作经历按 JSON 解析；申报等级按班级/人员规则。",
        "已解决",
    ),
]
for i, r in enumerate(rows, start=1):
    for j, val in enumerate(r):
        set_cell_text(
            t1.cell(i, j),
            val,
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT,
            fill="E2EFDA" if j == 3 else None,
        )
widths = [Cm(1.4), Cm(5.8), Cm(7.2), Cm(1.8)]
for row in t1.rows:
    for idx, cell in enumerate(row.cells):
        cell.width = widths[idx]

add_para("", space_after=6)

# 3
add_heading_cn("三、使用说明", 1)
add_heading_cn("3.1 入口位置", 2)
add_para("管理后台 → 开班准备 → 班级列表 → 打开「资料收集」→ 工具栏按钮：", size=12, space_after=4)
add_para(
    "①「默认模板导出」—— 生成整包材料（Word/图片/Excel），供认定归档使用。",
    size=12,
    space_after=2,
    first_line=0.74,
)
add_para(
    "②「仅导出名单」—— 仅导出《报名人员汇总表》Excel（25 列），不打包证件材料。",
    size=12,
    space_after=8,
    first_line=0.74,
)

add_heading_cn("3.2「默认模板导出」操作步骤", 2)
steps = [
    "确认班级已绑定资料收集计划，且人员资料已尽量填写完整（身份证、学历、合同三页、工作经历等）。",
    "在人员列表中可勾选部分人员导出；不勾选则导出当前计划下全部人员。",
    "点击「默认模板导出」，弹出进度窗口，显示「当前进度：已完成人数/总人数」。",
    "导出成功后点击「下载文件」，获得 ZIP 压缩包。",
    "解压后按文件夹核验材料（见 3.3）。",
]
for i, s in enumerate(steps, 1):
    add_para(f"{i}. {s}", size=12, space_after=3)

add_para("注意：", size=12, bold=True, space_before=6, space_after=2)
notes = [
    "导出为异步任务，人数较多时可能需 1～3 分钟，请勿重复连点；同计划同人员勾选会复用进行中任务。",
    "若进度提示「未找到压缩包」，系统会继续等待产物落盘/上传，无需立即判定失败。",
    "下载链接走平台文件服务，需登录态；请勿把链接外发到无权限环境。",
    "班级工种、等级决定认定申请表/评价登记表的默认「申报工种」「申报级别」；请先在班级资料中维护正确。",
]
for n in notes:
    add_para("• " + n, size=11, space_after=2)

add_heading_cn("3.3 导出包目录说明", 2)
t2 = doc.add_table(rows=1, cols=2)
set_table_borders(t2)
set_cell_text(t2.cell(0, 0), "文件夹/文件", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9E2F3")
set_cell_text(t2.cell(0, 1), "用途与核对要点", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9E2F3")
pack_rows = [
    (
        "工作年限证明/",
        "每人一份 Word。页边距应为上3.7、下3.5、左2.8、右2.6（厘米）；证明人/电话优先班级单位联系人。",
    ),
    (
        "职业技能评价登记表（2025新）/",
        "申报工种、申报级别与班级一致（无单独预约时）；有上一等级证书勾选「晋级评价」，否则「初次评价」。",
    ),
    (
        "职业技能等级认定申请表/",
        "毕业时间、企业电话/地址/邮编、工作经历、申报等级应已填充；现工种等级无材料时不应误填五级。",
    ),
    ("广西壮族自治区职业技能鉴定登记表/", "自治区鉴定登记表（随包生成）。"),
    ("小二寸证件照 / 身份证正面 / 身份证反面 / 身份证正反面合并/", "人员证件材料。"),
    ("合同 / 合同(合并) / 毕业证 / 上一等级证书 / 职称证/", "人员附件材料（有上传则导出）。"),
    ("*报名人员汇总表*.xlsx 等", "汇总 Excel 与认定相关列表。"),
]
for a, b in pack_rows:
    row = t2.add_row().cells
    set_cell_text(row[0], a, size=10)
    set_cell_text(row[1], b, size=10)

add_heading_cn("3.4「仅导出名单」说明（附）", 2)
add_para(
    "点击「仅导出名单」下载 Excel，文件名格式：{班级名称}-报名人员汇总表-{计划ID}.xlsx。"
    "表头为桌面《报名人员汇总表模板》A–Y 共 25 列（含「缺少的信息」）。勾选人员时仅导出勾选行。",
    first_line=0.74,
    size=12,
    space_after=8,
)

# 4
add_heading_cn("四、测试结果（验收用）", 1)
add_para(
    "下列用例对应《文档导出问题.docx》三项问题。建议使用已维护班级工种等级、且资料较完整的班级。"
    "「实测结果」栏由培训部门抽测后勾选。",
    first_line=0.74,
    size=12,
    space_after=6,
)

t3 = doc.add_table(rows=1, cols=5)
set_table_borders(t3)
for j, h in enumerate(["用例", "操作步骤", "预期结果", "实测结果", "备注"]):
    set_cell_text(t3.cell(0, j), h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9E2F3")

cases = [
    (
        "TC-01\n工作年限证明\n页边距",
        "1.资料收集点「默认模板导出」\n2.下载ZIP，打开「工作年限证明」中任一文稿\n3.Word→布局→页边距，查看数值",
        "上 3.7cm、下 3.5cm、左 2.8cm、右 2.6cm；正文排版整齐可读",
        "☐通过\n☐不通过",
        "",
    ),
    (
        "TC-02\n工作年限证明\n联系人",
        "班级已维护「单位联系人/电话」时导出工作年限证明",
        "文中证明人、联系电话优先为班级单位联系人信息",
        "☐通过\n☐不通过",
        "",
    ),
    (
        "TC-03\n评价登记表\n工种与级别",
        "选择班级等级非五级（如一级/二级）的班级导出；打开评价登记表",
        "申报工种与班级工种一致；申报级别与班级等级一致，不再一律「五级/初级工」",
        "☐通过\n☐不通过",
        "",
    ),
    (
        "TC-04\n评价登记表\n初评/晋级",
        "分别选取「有/无上一等级证书」人员导出登记表",
        "有证书：☑晋级评价；无证书：☑初次评价",
        "☐通过\n☐不通过",
        "",
    ),
    (
        "TC-05\n认定申请表\n字段完整性",
        "打开认定申请表，核对毕业时间、电话、地址、邮编、工作经历",
        "上述字段按人员/企业资料导出；源数据有值时不再空白",
        "☐通过\n☐不通过",
        "",
    ),
    (
        "TC-06\n认定申请表\n等级口径",
        "对照班级申报等级与人员现工种材料，检查申报等级与原/现等级",
        "申报等级与班级/预约规则一致；无现工种材料时不误填「五级」",
        "☐通过\n☐不通过",
        "",
    ),
    (
        "TC-07\n导出流程\n稳定性",
        "勾选多人导出至出现下载链接；对同一勾选重复点击",
        "能成功下载；不因短暂「未找到压缩包」误报失败；不重复开并行任务",
        "☐通过\n☐不通过",
        "",
    ),
    (
        "TC-08\n仅导出名单\n（附）",
        "点击「仅导出名单」，打开下载的 Excel",
        "文件名为「班级名称-报名人员汇总表-计划ID.xlsx」；含 25 列表头及人员数据",
        "☐通过\n☐不通过",
        "",
    ),
]
for c in cases:
    cells = t3.add_row().cells
    for j, val in enumerate(c):
        set_cell_text(
            cells[j],
            val,
            size=9,
            align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT,
        )

add_para("", space_after=6)
add_heading_cn("4.1 研发自测记录（供参考）", 2)
add_para(
    "研发侧已完成：工作年限证明模板页边距校验（3.7/3.5/2.8/2.6 cm）；申报等级回退班级等级单测；"
    "评价登记表晋级/初次勾选单测；认定申请表字段变量填充逻辑修复；默认模板导出异步任务与下载链路联调通过。"
    "最终业务验收以培训部门实测勾选为准。",
    first_line=0.74,
    size=11,
    space_after=8,
)

# 5
add_heading_cn("五、确认签字", 1)
add_para(
    "本人确认：已按本文件第三章操作完成抽测，第四章用例结果属实。"
    "对「默认模板导出」相关三项历史问题关闭无异议。",
    first_line=0.74,
    size=12,
    space_after=10,
)

sign = doc.add_table(rows=5, cols=4)
set_table_borders(sign)
sign_rows = [
    ("角色", "姓名（签字）", "日期", "意见"),
    ("培训部门确认人", "", "", "☐同意关闭问题\n☐需返工：________"),
    ("培训部门负责人", "", "", ""),
    ("技术对接人", "", "", ""),
    ("备注", "", "", ""),
]
for i, r in enumerate(sign_rows):
    for j, val in enumerate(r):
        set_cell_text(
            sign.cell(i, j),
            val,
            bold=(i == 0 or j == 0),
            size=10.5,
            align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT,
            fill="D9E2F3" if i == 0 else ("F5F5F5" if j == 0 else None),
        )
for cell in sign.rows[4].cells:
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(18)

add_para("", space_after=10)
add_para("（以下无正文）", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_para(
    "附件依据：《文档导出问题.docx》（企业微信，2026-09）。功能入口：开班资料收集工具栏「默认模板导出」。",
    size=9,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_before=6,
)

doc.save(out)
print("saved", out)
print("size", out.stat().st_size)
