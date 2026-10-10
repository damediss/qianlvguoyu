from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "04_requirements" / "contract-test-fixtures"

BODY_FONT = "宋体"
HEADING_FONT = "黑体"
BODY_SIZE = 10.5
CONTENT_WIDTH_DXA = 8900


def set_run_font(run, name=BODY_FONT, size=BODY_SIZE, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=BODY_SIZE):
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_font(paragraph.add_run(text), size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def set_table_borders(table, color="7F7F7F", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, before, after in (
        ("Heading 1", 12, 9, 4),
        ("Heading 2", 11, 6, 3),
        ("Heading 3", 10.5, 4, 2),
    ):
        style = doc.styles[style_name]
        style.font.name = HEADING_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), HEADING_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True


def configure_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Mm(297)
        section.page_height = Mm(210)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
    else:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(28)
        section.right_margin = Mm(25)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def set_header_footer(section, left_text, right_text="脱敏测试样本·仅用于系统验收"):
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(left_text), size=8.5, color="666666")
    set_run_font(paragraph.add_run("\t" + right_text), size=8.5, color="888888")
    tabs = paragraph.paragraph_format.tab_stops
    tabs.add_tab_stop(Mm(150))

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "BFBFBF")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer_p.add_run("第 "), size=8.5, color="777777")
    add_page_field(footer_p)
    set_run_font(footer_p.add_run(" 页"), size=8.5, color="777777")


def new_contract(title, contract_no, header_text):
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    set_header_footer(doc.sections[0], header_text)
    doc.core_properties.title = title
    doc.core_properties.subject = "WeKnora 历史 DOCX 派生 POC 脱敏源文档"
    doc.core_properties.author = "Guoyu Contract POC"
    doc.core_properties.comments = "全部主体、金额和项目均为虚构测试数据，不得用于签署。"

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(6)
    set_run_font(title_p.add_run(title), name=HEADING_FONT, size=18, bold=True)
    number_p = doc.add_paragraph()
    number_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    number_p.paragraph_format.space_after = Pt(12)
    set_run_font(number_p.add_run(f"合同编号：{contract_no}"), size=9.5, color="555555")
    return doc


def add_body_paragraph(doc, text, first_indent=True, bold_prefix=None, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(after)
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_field_paragraph(doc, label, value, split_value=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(label + "："), bold=True)
    if split_value and len(value) > 6:
        pivot = len(value) // 2
        set_run_font(p.add_run(value[:pivot]), size=11, italic=True)
        set_run_font(p.add_run(value[pivot:]), size=11, italic=True)
    else:
        set_run_font(p.add_run(value), size=11)
    return p


def add_clause(doc, title, paragraphs):
    doc.add_paragraph(title, style="Heading 1")
    for paragraph in paragraphs:
        add_body_paragraph(doc, paragraph)


def add_data_table(doc, headers, rows, widths, header_fill="E7E6E6"):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for index, header in enumerate(headers):
        shade_cell(table.rows[0].cells[index], header_fill)
        set_cell_text(table.rows[0].cells[index], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header_row(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[index], value, align=align)
    return table


def add_signature_table(doc, party_a, party_b):
    doc.add_paragraph("签署页", style="Heading 1")
    add_body_paragraph(doc, "（以下无正文，为本合同签署栏。双方确认已完整阅读并理解本合同全部条款。）")
    table = doc.add_table(rows=5, cols=2)
    set_table_geometry(table, [4450, 4450])
    set_table_borders(table, color="BFBFBF", size="4")
    left = [f"甲方（盖章）：{party_a}", "法定代表人/授权代表：", "联系人：张明（测试）", "联系电话：13800000001", "签署日期：2026年8月17日"]
    right = [f"乙方（盖章）：{party_b}", "法定代表人/授权代表：", "联系人：李华（测试）", "联系电话：13800000002", "签署日期：2026年8月17日"]
    for row, (a_text, b_text) in enumerate(zip(left, right)):
        set_cell_text(table.rows[row].cells[0], a_text)
        set_cell_text(table.rows[row].cells[1], b_text)


def build_framework_agreement(path):
    party_a = "广西启航产业发展有限公司"
    party_b = "南宁云枢数字科技有限公司"
    amount = "人民币壹佰捌拾万元（¥1,800,000.00）"
    project = "企业知识协同与合同管理平台建设项目"
    scope = "企业知识库、合同智能起草、流程集成、用户培训及一年期运维服务"
    doc = new_contract("软件技术服务框架协议", "GXQH-TECH-2026-0817", f"软件技术服务框架协议 | {party_a}")

    add_field_paragraph(doc, "甲方", party_a, split_value=True)
    add_field_paragraph(doc, "乙方", party_b, split_value=True)
    add_field_paragraph(doc, "项目名称", project)
    add_field_paragraph(doc, "合同金额", amount)
    add_field_paragraph(doc, "服务范围", scope)
    add_body_paragraph(doc, "鉴于甲方拟推进企业知识资产和合同业务数字化建设，乙方具备相关软件研发、实施和运维能力，双方依据《中华人民共和国民法典》等法律法规，在平等、自愿和诚实信用基础上订立本协议。")

    add_clause(doc, "第一条 合作目标与原则", [
        "双方围绕企业知识协同与合同管理平台开展长期合作，通过分阶段建设、持续优化和服务保障，提高甲方历史文件检索、合同复用、版本管理及审计追溯能力。",
        "本协议为框架性约定。具体需求、范围、排期和交付标准以双方确认的项目任务书、需求规格说明书或变更单为准。",
    ])
    add_clause(doc, "第二条 服务内容", [
        "乙方负责需求调研、总体方案设计、软件开发配置、系统部署、数据初始化、接口联调、试运行支持、用户培训和运维服务。",
        "涉及历史合同的处理应保留原始文件，语义检索结果仅用于选源和内容建议；新合同由经确认的源文档派生生成，并记录源文件哈希、结构清单版本和字段来源。",
    ])
    doc.add_paragraph("2.1 主要交付成果", style="Heading 2")
    add_data_table(doc, ["序号", "交付成果", "主要内容", "交付形式"], [
        ("1", "项目实施方案", "范围、组织、计划、风险和沟通机制", "DOCX/PDF"),
        ("2", "平台软件", "知识检索、合同派生、版本与权限模块", "部署包"),
        ("3", "技术文档", "需求、设计、接口、部署、运维说明", "DOCX/PDF"),
        ("4", "培训材料", "管理员及业务用户培训资料", "PPTX/DOCX"),
    ], [800, 1900, 4400, 1800])

    add_clause(doc, "第三条 项目计划与变更", [
        "项目计划分为启动与调研、方案与设计、开发与配置、联调与试运行、验收与运维五个阶段。双方项目负责人每周同步进展，并对影响范围、费用或工期的事项形成书面记录。",
        "任何一方提出需求变更时，应说明变更原因、内容和优先级。乙方完成影响评估后，由双方授权代表确认变更单；未确认前仍执行原计划。",
    ])
    add_clause(doc, "第四条 合同价款与支付", [
        f"本合同总价为{amount}，该价格已包含乙方履行本合同所需的软件开发、实施、培训、差旅、税费及约定期限内运维服务费用。",
    ])
    add_data_table(doc, ["阶段", "付款比例", "付款条件", "金额"], [
        ("首付款", "30%", "合同生效且乙方提交项目实施方案后", "¥540,000.00"),
        ("里程碑款", "40%", "核心功能完成并进入试运行后", "¥720,000.00"),
        ("验收款", "25%", "项目终验通过后", "¥450,000.00"),
        ("质保尾款", "5%", "终验满三个月且无重大遗留问题", "¥90,000.00"),
    ], [1400, 1200, 4100, 2200])

    add_clause(doc, "第五条 交付与验收", [
        "乙方按计划提交交付成果，甲方应在收到后十个工作日内组织验收并书面反馈。验收以双方确认的需求、测试用例、性能指标和安全要求为依据。",
        "发现缺陷时，乙方应在双方约定期限内整改并重新提交。甲方无正当理由逾期未反馈的，不视为自动验收，双方项目负责人应及时升级协调。",
    ])
    add_clause(doc, "第六条 知识产权与数据", [
        "双方在本合同签署前已拥有的技术、软件、文档和其他知识产权仍归原权利人所有。乙方通用产品和基础组件的知识产权归乙方所有。",
        "甲方业务数据、历史合同、知识文档及其衍生数据归甲方所有。未经甲方书面许可，乙方不得用于本项目以外的训练、测试、展示或其他用途。",
    ])
    add_clause(doc, "第七条 保密与安全", [
        "双方对履约过程中知悉的商业秘密、个人信息、合同文件、系统配置和安全资料承担保密义务。除履约所必需人员外，不得向第三方披露。",
        "乙方应落实最小权限、租户隔离、操作审计、传输加密、私有存储和安全备份。发生安全事件时，应立即采取措施并在二十四小时内通知甲方。",
    ])
    add_clause(doc, "第八条 运维与服务保障", [
        "终验通过后，乙方提供十二个月免费运维服务，包括故障处理、版本修复、运行监控和技术咨询。重大故障响应时间不超过三十分钟，一般故障响应时间不超过两小时。",
    ])
    add_clause(doc, "第九条 违约责任与合同解除", [
        "任何一方违反合同并造成对方损失的，应承担相应赔偿责任。一方严重违约且在收到书面通知后十五个工作日内未改正的，守约方有权解除合同。",
    ])
    add_clause(doc, "第十条 争议解决与其他", [
        "本合同适用中华人民共和国法律。履行过程中发生争议，双方应友好协商；协商不成的，任一方可向甲方所在地有管辖权的人民法院起诉。",
        "本合同一式肆份，双方各执贰份，自双方盖章之日起生效。附件、任务书和经确认的变更单均为本合同组成部分，与本合同具有同等法律效力。",
    ])
    add_signature_table(doc, party_a, party_b)
    doc.save(path)


def build_development_contract(path):
    party_a = "桂林智造装备有限公司"
    party_b = "广西星河软件技术有限公司"
    amount = "人民币玖拾陆万元（¥960,000.00）"
    project = "生产运营数字化管理系统开发实施项目"
    scope = "生产计划、设备台账、质量追溯、经营看板及既有ERP接口集成"
    doc = new_contract("软件系统开发实施合同", "GLZZ-SW-2026-036", "软件系统开发实施合同 | 项目资料")

    source_table = doc.add_table(rows=5, cols=2)
    set_table_geometry(source_table, [2100, 6800])
    set_table_borders(source_table)
    for row, (label, value) in enumerate((
        ("甲方：", party_a),
        ("乙方：", party_b),
        ("项目名称：", project),
        ("合同金额：", amount),
        ("服务范围：", scope),
    )):
        shade_cell(source_table.rows[row].cells[0], "E7E6E6")
        set_cell_text(source_table.rows[row].cells[0], label, bold=True)
        set_cell_text(source_table.rows[row].cells[1], value)

    add_body_paragraph(doc, "甲乙双方就生产运营数字化管理系统的定制开发、部署实施和试运行服务达成如下约定。合同所载公司、联系人、金额和项目均为脱敏测试数据。")
    add_clause(doc, "第一条 项目范围", [
        "乙方依据经确认的需求规格说明书完成系统设计、编码、测试、部署和培训。系统包含生产计划、设备管理、质量追溯、经营分析和系统管理等模块。",
        "与甲方现有ERP、统一身份认证和消息平台的接口，由双方共同确定数据字段、调用频率、异常处理和安全策略。",
    ])
    add_clause(doc, "第二条 里程碑与交付物", [
        "项目总工期为一百二十个自然日，自项目启动会召开之日起计算。因甲方需求确认、环境准备或第三方接口原因造成的延误，双方据实调整计划。",
    ])
    add_data_table(doc, ["里程碑", "计划周期", "主要工作", "交付/确认材料"], [
        ("M1 项目启动", "第1-10日", "调研、范围确认、项目计划", "实施计划、调研纪要"),
        ("M2 方案设计", "第11-30日", "原型、架构、接口和数据设计", "需求规格、概要设计"),
        ("M3 开发测试", "第31-80日", "功能开发、单元测试、集成测试", "测试版本、测试报告"),
        ("M4 上线试运行", "第81-105日", "部署、培训、数据初始化", "部署记录、培训签到"),
        ("M5 项目验收", "第106-120日", "问题整改、终验和移交", "验收报告、运维手册"),
    ], [1700, 1400, 3300, 2500])

    add_clause(doc, "第三条 费用与结算", [
        f"合同金额为{amount}，其中软件开发实施费人民币捌拾陆万元，第三方测试及安全测评配合费人民币壹拾万元。",
    ])
    add_data_table(doc, ["付款节点", "比例", "触发条件", "含税金额"], [
        ("合同生效", "30%", "合同签署且收到合规发票", "¥288,000.00"),
        ("设计确认", "20%", "需求规格和概要设计确认", "¥192,000.00"),
        ("上线试运行", "30%", "系统部署并进入试运行", "¥288,000.00"),
        ("终验通过", "20%", "双方签署终验报告", "¥192,000.00"),
    ], [1700, 1100, 3900, 2200])

    add_clause(doc, "第四条 双方职责", [
        "甲方负责明确业务需求、提供测试数据和部署环境、组织关键用户参与评审，并按期完成确认及付款。",
        "乙方负责安排合格项目团队，遵循甲方安全管理制度，按计划交付成果，对开发质量和交付完整性负责。",
    ])
    add_clause(doc, "第五条 验收标准", [
        "功能验收以需求规格说明书为依据；性能、安全和兼容性指标以测试方案为依据。严重缺陷清零、一般缺陷具有双方认可的处理计划后方可通过终验。",
        "乙方应提交源代码、可执行程序、数据库脚本、部署说明、接口文档、测试报告和管理员手册，具体范围以附件交付清单为准。",
    ])
    add_clause(doc, "第六条 知识产权", [
        "甲方支付全部合同款后，针对甲方业务需求形成的定制代码和项目文档，其使用权归甲方所有；乙方既有平台、通用框架、工具和第三方开源组件的权利归原权利人所有。",
        "乙方保证交付成果不侵犯第三方合法权益。如发生权利主张，乙方应及时处理并承担因其侵权行为造成的直接损失。",
    ])
    add_clause(doc, "第七条 质量保证与运维", [
        "项目终验后提供十二个月质量保证。质保期内因乙方交付缺陷造成的问题由乙方免费修复；新增需求另行评估。",
    ])
    add_clause(doc, "第八条 保密、违约与争议", [
        "双方对业务数据、源代码、技术资料和商务信息承担保密义务。未经授权不得复制、传播或用于本合同以外的目的。",
        "争议由双方协商解决；协商不成的，向甲方所在地有管辖权的人民法院起诉。",
    ])
    add_signature_table(doc, party_a, party_b)
    doc.save(path)


def build_operations_contract(path):
    party_a = "柳州数联公共服务有限公司"
    party_b = "广西安维信息技术有限公司"
    amount = "人民币陆拾万元（¥600,000.00）"
    project = "企业信息平台年度运行维护服务项目"
    scope = "应用系统巡检、故障响应、安全加固、数据备份、版本维护及季度运营报告"
    doc = new_contract("信息平台运行维护服务合同", "LZSL-OPS-2026-017", f"运行维护服务合同 | {party_a}")

    add_field_paragraph(doc, "甲方", party_a)
    add_field_paragraph(doc, "乙方", party_b)
    add_field_paragraph(doc, "合同金额", amount)
    add_field_paragraph(doc, "项目名称", project)
    add_field_paragraph(doc, "服务范围", scope)
    add_body_paragraph(doc, "双方就甲方现有企业信息平台的年度运行维护、故障处置、安全保障和持续优化服务达成本合同。服务期为2026年9月1日至2027年8月31日。")

    add_clause(doc, "第一条 服务目标", [
        "乙方通过常态巡检、分级响应、备份恢复演练和安全加固，保障平台可用性、数据完整性和业务连续性，并形成可审计的服务记录。",
    ])
    add_clause(doc, "第二条 服务内容", [
        "服务包括应用系统、数据库、中间件和接口运行检查；故障定位与恢复；补丁和版本维护；备份检查与恢复演练；安全事件协同处置；用户咨询和季度运行分析。",
        "超出本合同约定范围的重大功能开发、基础设施采购和第三方许可费用，由双方另行签署任务书或补充协议。",
    ])
    add_clause(doc, "第三条 服务等级", [
        "乙方设置7×24小时紧急联系人。事件等级以业务影响、用户范围、持续时间和数据风险综合确定。",
    ])
    add_data_table(doc, ["等级", "典型情形", "响应时限", "恢复/方案时限"], [
        ("P1 紧急", "核心业务中断、重大数据或安全风险", "15分钟", "2小时恢复或提供应急方案"),
        ("P2 高", "重要功能不可用且无替代方案", "30分钟", "4小时恢复或提供绕行方案"),
        ("P3 中", "部分功能异常，对业务影响有限", "2小时", "1个工作日内处理"),
        ("P4 低", "咨询、配置优化和一般问题", "4小时", "3个工作日内处理"),
    ], [1200, 3900, 1500, 2300])

    add_clause(doc, "第四条 服务过程与报告", [
        "乙方应使用服务工单记录问题来源、影响范围、处理过程、结果、责任人和时间。涉及数据修改、权限调整和生产发布的操作须履行甲方审批流程。",
        "乙方每月提交服务统计，每季度提交运行分析报告，内容包括可用性、事件分布、问题趋势、容量变化、安全风险和改进建议。",
    ])
    add_clause(doc, "第五条 服务费用与支付", [
        f"本合同服务费为{amount}，按季度支付。每季度服务完成且甲方确认季度报告后，乙方开具合法有效的增值税专用发票。",
    ])
    add_data_table(doc, ["期次", "服务期间", "付款比例", "金额"], [
        ("第一期", "2026.09-2026.11", "25%", "¥150,000.00"),
        ("第二期", "2026.12-2027.02", "25%", "¥150,000.00"),
        ("第三期", "2027.03-2027.05", "25%", "¥150,000.00"),
        ("第四期", "2027.06-2027.08", "25%", "¥150,000.00"),
    ], [1500, 2600, 1600, 3200])

    add_clause(doc, "第六条 数据安全与保密", [
        "乙方人员仅可在授权范围内访问甲方系统和数据，不得共享账号，不得将生产数据复制至未经批准的设备或环境。远程访问应使用甲方认可的安全通道并保留审计日志。",
        "合同终止后，乙方应归还或安全删除甲方资料，并书面确认删除结果；法律法规另有保存要求的除外。",
    ])
    add_clause(doc, "第七条 考核与违约", [
        "甲方按季度对响应及时性、问题关闭率、报告质量和用户满意度进行考核。因乙方原因导致连续或重大服务中断的，甲方有权要求整改、扣减相应服务费并追究损失。",
    ])
    add_clause(doc, "第八条 合同终止与交接", [
        "合同期满或提前终止时，乙方应完成账号、配置、文档、未结工单和风险事项交接，配合甲方或后续服务方平稳接管。",
    ])
    add_clause(doc, "第九条 争议解决", [
        "本合同未尽事宜由双方协商并签署补充协议。争议协商不成的，向甲方所在地有管辖权的人民法院起诉。",
    ])

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=True)
    set_header_footer(section, f"附件一 服务资产与巡检清单 | {party_b}")
    doc.add_paragraph("附件一 服务资产与巡检清单", style="Heading 1")
    add_body_paragraph(doc, "本附件用于验证多 section、横向页面和宽表格结构在源文档派生过程中保持不变。")
    add_data_table(doc, ["资产类别", "示例对象", "巡检频率", "检查重点", "记录"], [
        ("应用系统", "合同管理、知识服务、统一门户", "每日", "可用性、任务积压、接口错误", "巡检工单"),
        ("数据库", "MySQL主从集群", "每日", "连接、慢查询、复制延迟、备份", "数据库巡检表"),
        ("中间件", "网关、缓存、消息队列", "每日", "容量、连接、堆积、异常日志", "中间件巡检表"),
        ("安全设施", "WAF、主机防护、审计平台", "每周", "告警、规则、漏洞和处置状态", "安全周报"),
        ("备份系统", "数据库备份、对象存储备份", "每周", "任务成功率、保留周期、抽样恢复", "备份核查记录"),
    ], [1500, 2600, 1400, 2800, 1600])

    final_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(final_section)
    set_header_footer(final_section, f"运行维护服务合同 | {party_a} / {party_b}")
    add_signature_table(doc, party_a, party_b)
    doc.save(path)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = [
        OUTPUT_DIR / "01_软件技术服务框架协议_源文档.docx",
        OUTPUT_DIR / "02_软件系统开发实施合同_源文档.docx",
        OUTPUT_DIR / "03_信息平台运维服务合同_源文档.docx",
    ]
    build_framework_agreement(outputs[0])
    build_development_contract(outputs[1])
    build_operations_contract(outputs[2])
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
