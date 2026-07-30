from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "/Users/macmain/Documents/qianlvguoyu/三个月数字化与AI建设计划.docx"


BLUE = (31, 78, 121)
LIGHT_BLUE = "EAF2F8"
HEADER_BLUE = "2F5597"
GRAY = (89, 89, 89)
TEXT = (31, 31, 31)


def cn_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        child = tc_mar.find(qn(f"w:{name}"))
        if child is None:
            child = OxmlElement(f"w:{name}")
            tc_mar.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")


def set_cell_width(cell, cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def dont_split_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    keep_with_next(p)
    r = p.add_run(text)
    cn_font(r, 15 if level == 1 else 12.5, True, BLUE)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        cn_font(r, 10.5, True, TEXT)
        r = p.add_run(text[len(bold_prefix):])
        cn_font(r, 10.5, False, TEXT)
    else:
        r = p.add_run(text)
        cn_font(r, 10.5, False, TEXT)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(item)
        cn_font(r, 10.3, False, TEXT)


def add_table(doc, headers, rows, widths, font_size=8.6, center_cols=None):
    center_cols = center_cols or set()
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for ri, row in enumerate(table.rows):
        dont_split_row(row)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ci < len(widths):
                set_cell_width(cell, widths[ci])
            if ri == 0:
                set_shading(cell, HEADER_BLUE)
            elif ri % 2 == 0:
                set_shading(cell, "F7F9FC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.06
                if ri == 0 or ci in center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    cn_font(r, font_size, ri == 0, (255, 255, 255) if ri == 0 else TEXT)
    repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, 130, 160, 130, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    cn_font(r, 11, True, BLUE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(body)
    cn_font(r, 10.3, False, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.65)
section.right_margin = Cm(1.65)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
r = title.add_run("三个月数字化与 AI 建设计划")
cn_font(r, 24, True, BLUE)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(2)
r = subtitle.add_run("基于现有调研资料、系统进度表、培训平台原型与资源预算方向")
cn_font(r, 11.5, False, GRAY)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(14)
r = meta.add_run("版本：初稿 v1.0｜日期：2026年4月29日｜周期：90天")
cn_font(r, 10, False, GRAY)

add_callout(
    doc,
    "总体定位",
    "未来三个月不做泛化的“AI 展示”，重点围绕培训、竞赛、课程、知识库和项目交付建立可使用、可验收、可持续迭代的数字化底座。"
)

add_heading(doc, "一、建设背景与当前判断")
add_para(
    doc,
    "现有资料显示，公司当前同时推进企业培训中心 SaaS 云平台、企业竞赛系统、公司智能体、课程初稿开发工作流、AI 精准学、服务器选型和本地模型部署等事项。整体特点是目标清晰但时间紧、外包与自建并行、业务资料分散、需求文档和验收标准仍需补齐。"
)
add_bullets(
    doc,
    [
        "广西千铝工投教育科技发展有限公司是当前培训、认定、课程、考试等业务数字化的主场景。",
        "广西国昱数字科技有限公司承担竞赛平台、企业数字化项目、智能体和本地模型部署等技术支撑方向。",
        "广西国之重器工程科技有限公司的具体业务、客户和项目仍需在第一阶段补充调研，暂纳入集团统一数字资产和项目治理范围。",
        "三个月内的重点不是一次性做完整系统，而是用 MVP 跑通主流程，并用项目治理机制控制外包进度、源码资料、服务器权限和验收边界。"
    ]
)

add_heading(doc, "二、三个月总目标")
add_table(
    doc,
    ["目标", "三个月内形成的结果", "验收口径"],
    [
        ["一个数字化底座", "统一项目台账、资料归档、代码与外包成果交接、知识库目录、基础服务器/存储方案", "核心资料有目录、有负责人、有权限、有备份"],
        ["两个核心业务系统", "培训 SaaS 一期可用；竞赛系统支持企业竞赛和政府全区竞赛阶段使用", "报名、学习/赛事、题库、考试/评分、成绩、证书/排名、导出等核心流程可跑通"],
        ["三类 AI 能力", "公司智能体、课程初稿工作流、培训 AI 精准学形成可试用版本", "至少完成首批部门试点、1门课程样板、错题/薄弱点/推荐练习 MVP"],
        ["一套治理机制", "形成周报、里程碑、风险、外包交付、验收标准和数据安全机制", "每周可向老板汇报进度、风险、需协调事项和下周计划"],
    ],
    [4.0, 8.3, 5.0],
    8.8,
    {0, 2},
)

add_heading(doc, "三、推进原则")
add_table(
    doc,
    ["原则", "要求"],
    [
        ["先保交付，再扩功能", "培训 SaaS 和竞赛系统优先满足对外承诺和业务主流程，复杂报表、深度算法和精细化运营功能后置。"],
        ["先定验收，再谈开发", "所有外包开发必须先形成 MVP 功能清单、角色权限、数据样例、验收标准和交付物清单。"],
        ["先云端/混合验证，再本地重投入", "本地模型部署先完成必要性评估、数据安全要求、成本测算和维护责任，再决定采购规模。"],
        ["AI 围绕业务资料落地", "智能体和课程工作流必须绑定资料库、流程、模板和人工审核，不做脱离业务的演示型功能。"],
        ["源码、数据、账号归公司", "外包成果必须交付源码、数据库结构、接口文档、部署文档、账号权限和维护说明。"],
    ],
    [4.0, 13.3],
    9.2,
    {0},
)

add_heading(doc, "四、重点项目群")
add_table(
    doc,
    ["项目群", "核心内容", "优先级", "关键交付物"],
    [
        ["培训 SaaS 云平台", "企业账号、学员管理、课程学习、报名、题库、考试、成绩、证书/学习证明、后台统计、AI 精准学入口", "P0", "一期 MVP 需求书、原型、测试版本、上线验收表"],
        ["企业竞赛系统", "报名、赛程、题库、考试/评分、排名、证书或结果公示、数据导出、活动阶段上线", "P0", "竞赛流程需求、阶段上线计划、压测与应急预案"],
        ["公司智能体与知识库", "首批部门资料喂养、FAQ、制度流程、项目资料、客户/培训问答、内部知识检索", "P1", "部门知识库目录、智能体试点、反馈表"],
        ["课程 AI 开发工作流", "资料输入、大纲生成、课件初稿、讲义、题库、短视频脚本、人工审核模板", "P1", "1门样板课程、提示词模板、审核标准"],
        ["培训 AI 精准学", "错题记录、薄弱知识点、推荐练习、学习报告、题库标签与知识点映射", "P1", "精准学 MVP 方案、数据字段、平台入口"],
        ["基础设施与数据安全", "NAS/备份、AI 工具、VPN、服务器选型、本地/云端/混合模型方案", "P2", "资源方案、预算引用、安全说明"],
        ["项目治理与外包管理", "项目看板、周报、需求确认、开发排期、测试验收、源码文档交接", "P0", "系统进度表、外包跟踪表、验收模板"],
    ],
    [3.2, 7.1, 1.9, 5.1],
    8.2,
    {0, 2},
)

add_heading(doc, "五、90天阶段计划")
add_table(
    doc,
    ["阶段", "时间", "主线任务", "阶段成果"],
    [
        ["第一阶段：摸清现状与快速样板", "第1-2周", "补齐三家公司业务、人员、系统、项目和资料图谱；确定培训 SaaS MVP；建立项目看板；上线公司智能体雏形；跑通课程初稿样板；完成服务器选型方案初版。", "老板可看到优先级清单、项目台账、智能体雏形和课程 AI 样板。"],
        ["第二阶段：培训平台一期上线", "第3-4周", "锁定技术路线和外包边界；完成培训主流程需求、角色权限、原型和数据样例；开发/改造培训 SaaS 一期；接入 AI 精准学基础入口。", "培训报名、课程、题库、考试、成绩、证书/证明等主流程可测试或试运行。"],
        ["第三阶段：竞赛系统与部门智能体扩展", "第5-8周", "推进竞赛系统阶段上线；完成报名、题库、评分、排名和导出；扩展各部门知识库与智能体；沉淀课程工作流模板；建立运维与应急机制。", "竞赛系统可支撑企业竞赛和政府竞赛启动，部门智能体完成首批配置。"],
        ["第四阶段：稳定运行与内化能力", "第9-12周", "优化平台稳定性、数据报表、权限和审计；完成外包交付验收；沉淀标准化 SOP；评估本地模型部署与二期建设清单。", "形成三个月复盘、二期路线图、可持续迭代机制和内部运维交接。"],
    ],
    [3.1, 2.5, 8.2, 3.5],
    8.6,
    {0, 1},
)

add_heading(doc, "六、周度 WBS")
add_table(
    doc,
    ["周次", "重点工作", "输出物"],
    [
        ["W1", "访谈老板、培训、课程、认定、竞赛、技术和外包相关人员；梳理三家公司业务、人员、系统和项目清单。", "业务图、人员图、系统流程图、问题优先级表"],
        ["W2", "确定培训 SaaS MVP、外包流程、智能体首批部门、课程样板、服务器选型初稿。", "MVP 清单、验收标准、智能体雏形、课程 AI 样板、服务器方案初稿"],
        ["W3", "完成培训平台原型确认、角色权限、数据样例；明确 Node.js 新开发与 Java 复用路线。", "需求文档、原型确认、技术路线对比、开发排期"],
        ["W4", "培训 SaaS 一期开发/改造与测试；AI 精准学入口设计；题库和考试模块联调。", "可测试版本、测试问题清单、精准学 MVP 设计"],
        ["W5", "竞赛系统需求锁定；拆分政府全区竞赛阶段上线范围；建立竞赛测试和应急方案。", "竞赛需求书、阶段上线表、风险清单"],
        ["W6", "竞赛报名、题库、考试/评分、排名、导出等核心模块开发和联调。", "竞赛系统测试版、接口/数据清单"],
        ["W7", "部门智能体扩展；课程工作流标准化；培训平台试运行问题修复。", "部门知识库、提示词模板、平台修复清单"],
        ["W8", "竞赛系统阶段上线演练；培训平台一期验收；外包阶段交付检查。", "阶段上线报告、验收表、源码/文档交接清单"],
        ["W9", "补齐报表、权限、审计、备份和日志；完善 AI 精准学学习报告。", "报表样例、权限表、备份策略、学习报告样板"],
        ["W10", "推进本地/云端/混合模型部署评估；形成数据安全和运维要求。", "部署对比表、配置清单、安全说明"],
        ["W11", "固化培训、竞赛、课程、智能体四类 SOP；组织内部使用培训。", "SOP 文档、培训记录、问题反馈表"],
        ["W12", "三个月复盘、二期路线图、预算复核、团队分工和长期运维机制。", "复盘报告、二期计划、运维交接表"],
    ],
    [1.4, 10.2, 5.7],
    8.0,
    {0},
)

add_heading(doc, "七、关键里程碑")
add_table(
    doc,
    ["节点", "时间", "必须完成", "老板可验收的结果"],
    [
        ["M1", "第14天", "智能体雏形、课程初稿样板、培训 SaaS MVP、服务器方案初稿、项目看板", "能看到 AI 和项目治理的可见成果"],
        ["M2", "第30天", "培训 SaaS 一期可用，AI 精准学基础入口完成", "培训主流程可跑通，可进入试运行"],
        ["M3", "第60天", "竞赛系统阶段上线，首批部门智能体完成配置", "可支撑企业竞赛和政府全区竞赛启动"],
        ["M4", "第90天", "平台稳定运行、外包交付验收、二期路线图和运维机制完成", "从项目推动转为持续运营和迭代"],
    ],
    [1.7, 2.4, 8.1, 5.1],
    8.6,
    {0, 1},
)

add_heading(doc, "八、AI 建设落地路径")
add_table(
    doc,
    ["AI 场景", "第一步", "三个月目标", "人工把关"],
    [
        ["公司智能体", "先收集制度、流程、项目资料、FAQ 和常用模板", "完成首批部门知识库和问答试用", "部门负责人确认资料范围和答案准确性"],
        ["课程初稿工作流", "选 1 门课程做样板，跑通资料输入到课件初稿和题库", "形成课程开发模板库，逐步替代人工初稿", "课程负责人审核专业性、合规性和表达质量"],
        ["AI 精准学", "定义错题、知识点、推荐练习、学习报告的数据字段", "在培训平台形成基础个性化学习能力", "教务/课程人员维护题库标签和知识点映射"],
        ["项目管理助手", "把项目进度、风险、周报模板结构化", "自动辅助生成周报、风险提醒和会议纪要", "项目负责人确认风险和节点变化"],
    ],
    [3.0, 5.1, 5.3, 3.9],
    8.3,
    {0},
)

add_heading(doc, "九、项目治理机制")
add_bullets(
    doc,
    [
        "每周固定输出一版项目周报：本周完成、下周计划、当前风险、需要老板协调事项。",
        "所有系统项目必须有五件套：需求清单、原型/流程、排期、验收标准、交付物清单。",
        "外包会议后 24 小时内形成纪要，明确负责人、截止时间、交付标准和未决问题。",
        "项目看板按 P0/P1/P2 管理，P0 项目每天跟进，P1 每周推进，P2 按节点复核。",
        "源码、数据库、接口、服务器账号、部署文档和运维说明必须纳入交付验收，不得只验页面。"
    ]
)

add_heading(doc, "十、风险与应对")
add_table(
    doc,
    ["风险", "影响", "应对动作"],
    [
        ["培训平台技术路线未定", "影响一个月上线和外包报价", "一周内完成 Node.js 新开发与 Java 平台复用对比，老板拍板后锁定范围"],
        ["需求只有大纲", "容易返工和扯皮", "先补 MVP 功能清单、数据样例、角色权限和验收标准"],
        ["竞赛活动时间固定", "系统稳定性和上线风险高", "拆分阶段上线，提前压测，准备人工兜底流程和应急联系人"],
        ["资料分散且质量不一", "影响智能体和课程 AI 输出质量", "建立资料目录、命名规范、负责人和审核机制"],
        ["本地模型投入较大", "采购后维护压力和成本风险高", "先做云端/本地/混合对比，明确必要性、预算和维护责任后采购"],
        ["外包成果不可控", "后续无法维护或二次开发", "合同与验收中写明源码、文档、部署、账号、数据库和维护交接"],
    ],
    [4.0, 4.3, 9.0],
    8.5,
    {0},
)

add_heading(doc, "十一、需要老板拍板事项")
add_table(
    doc,
    ["事项", "建议拍板时间", "需要明确"],
    [
        ["三个月优先级", "第1周", "培训 SaaS、竞赛系统、AI 智能体、课程工作流、服务器采购之间的优先顺序"],
        ["培训平台技术路线", "第1-2周", "Node.js 新开发还是 Java 平台复用，或采用阶段性混合方案"],
        ["外包主责与内部确认人", "第1-2周", "谁写需求、谁确认范围、谁验收、谁保管源码和账号"],
        ["服务器与本地模型预算边界", "第2-4周", "是否先采购 AI 电脑服务器，预算上限、部署方式和维护责任"],
        ["跨部门协调授权", "第1周", "哪些人员必须配合资料、需求、测试、验收和智能体喂养"],
    ],
    [5.2, 3.0, 9.1],
    8.8,
    {1},
)

add_heading(doc, "十二、三个月结束时的交付清单")
add_table(
    doc,
    ["类别", "交付物"],
    [
        ["系统交付", "培训 SaaS 一期、竞赛系统阶段上线版本、AI 精准学入口和学习报告样板"],
        ["AI 交付", "公司智能体首批试点、部门知识库目录、课程 AI 工作流模板、1门以上样板课程"],
        ["管理交付", "项目看板、周报模板、外包跟踪表、需求和验收模板、风险清单"],
        ["数据交付", "课程、题库、证书、项目、客户/企业、人才数据等核心数据目录和样例"],
        ["运维交付", "源码/数据库/接口/部署文档、账号权限清单、备份方案、二期路线图"],
    ],
    [3.2, 14.1],
    9.0,
    {0},
)

add_heading(doc, "附：第一周立即启动清单")
add_table(
    doc,
    ["序号", "动作", "完成标准"],
    [
        ["1", "确认三家公司业务、人员、项目、系统负责人", "形成一页负责人清单"],
        ["2", "确认培训 SaaS 一个月内必须上线功能", "形成 P0 功能清单和延后清单"],
        ["3", "组织 Node.js 与 Java 外包路线对比", "形成周期、报价、源码、维护、风险对比表"],
        ["4", "建立项目周报和看板", "所有 P0/P1/P2 项目进入跟踪表"],
        ["5", "选定首批智能体试点部门和资料清单", "资料开始喂养，2周内可试用"],
        ["6", "选 1 门课程做 AI 初稿样板", "输出大纲、课件初稿、题库和审核表"],
        ["7", "启动服务器和本地模型选型", "形成云端、本地、混合部署对比"],
    ],
    [1.3, 8.0, 8.0],
    8.8,
    {0},
)

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("三个月数字化与 AI 建设计划｜内部讨论稿")
cn_font(r, 8.5, False, GRAY)

doc.save(OUT)
print(OUT)
