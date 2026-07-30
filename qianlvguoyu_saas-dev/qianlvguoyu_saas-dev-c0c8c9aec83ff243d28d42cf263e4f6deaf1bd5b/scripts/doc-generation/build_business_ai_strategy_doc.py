from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("/Users/macmain/Documents/qianlvguoyu")
SRC = ROOT / "02_草稿与过程文档" / "公司业务AI化建设思路-草稿.md"
OUT = ROOT / "01_正式交付文档" / "公司业务三个月AI改造行动计划.docx"

BLUE = (31, 78, 121)
DARK = (31, 31, 31)
GRAY = (89, 89, 89)
HEADER_FILL = "2F5597"
ALT_FILL = "F7F9FC"
CALLOUT_FILL = "EAF2F8"


def cn_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        child = tc_mar.find(qn(f"w:{name}"))
        if child is None:
            child = OxmlElement(f"w:{name}")
            tc_mar.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")


def set_cell_width(cell, cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def dont_split_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_title_page(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(170)
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run("公司业务三个月 AI 改造行动计划")
    cn_font(r, 26, True, BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    r = subtitle.add_run("收回外包｜自研自开发｜统一技术栈｜平台 AI 化｜Agent 代理化")
    cn_font(r, 12, False, GRAY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(28)
    r = meta.add_run("版本：草稿版 v0.1｜日期：2026 年 4 月 30 日")
    cn_font(r, 10.5, False, GRAY)

    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    set_shading(cell, CALLOUT_FILL)
    set_cell_margins(cell, 160, 180, 160, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("三个月后要完成的结果：完成 Agent 面板第一版、公司内部知识库第一版、内部知识库与业务流程搭建、全员 AI 能力接入和平台 AI/Agent 接入试点；让公司从“人盯流程”变成“流程推着人走”，AI 负责生成、检查、提醒、汇总；完成外包项目收回清单和第一批交付接管；完成各业务 AI 化试点、2-5 个 Agent 试点闭环、旧 PHP 系统 Java 重构路线、竞赛系统重构范围、工业仿真收回样板和开发流程重构机制。")
    cn_font(r, 11, True, BLUE)

    doc.add_page_break()


def add_static_toc(doc, lines):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("目录")
    cn_font(r, 17, True, BLUE)

    for line in lines:
        if line.startswith("## "):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            cn_font(r, 10.5, True, DARK)
        elif line.startswith("### "):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(text)
            cn_font(r, 9.2, False, GRAY)
    doc.add_page_break()


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.style = f"Heading {min(level, 4)}"
    p.paragraph_format.space_before = Pt(8 if level <= 2 else 5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    size = {1: 22, 2: 15, 3: 12.3, 4: 10.8}.get(level, 10.5)
    color = BLUE if level <= 3 else DARK
    r = p.add_run(text)
    cn_font(r, size, True, color)
    return p


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    cn_font(r, 10.2, False, DARK)
    return p


def add_list_item(doc, text, ordered=False):
    p = doc.add_paragraph(style=None if ordered else "List Bullet")
    if ordered:
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.15)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    cn_font(r, 10, False, DARK)
    return p


def add_table(doc, rows):
    if not rows:
        return
    headers = rows[0]
    body = rows[1:]
    col_count = len(headers)
    table = doc.add_table(rows=1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in body:
        cells = table.add_row().cells
        for i in range(col_count):
            cells[i].text = row[i] if i < len(row) else ""

    page_width = 17.8
    if col_count == 2:
        widths = [4.2, page_width - 4.2]
    elif col_count == 3:
        widths = [3.9, 5.3, 8.6]
    elif col_count == 4:
        widths = [3.5, 4.8, 4.3, 5.2]
    else:
        widths = [page_width / col_count] * col_count

    for ri, row in enumerate(table.rows):
        dont_split_row(row)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ci < len(widths):
                set_cell_width(cell, widths[ci])
            if ri == 0:
                set_shading(cell, HEADER_FILL)
            elif ri % 2 == 0:
                set_shading(cell, ALT_FILL)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if ri == 0 or (ci == 0 and col_count <= 3):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    cn_font(r, 8.5, ri == 0, (255, 255, 255) if ri == 0 else DARK)
    repeat_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build_doc():
    source_lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.2)

    add_title_page(doc)
    add_static_toc(doc, [line for line in source_lines if line.startswith("## ") or line.startswith("### ")])

    i = 0
    while i < len(source_lines):
        line = source_lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# "):
            # Title is already represented by the cover page.
            i += 1
            continue
        if stripped.startswith("#### "):
            add_heading(doc, stripped[5:].strip(), 4)
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:].strip(), 3)
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, stripped[3:].strip(), 2)
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(source_lines, i)
            add_table(doc, rows)
            continue
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            add_list_item(doc, stripped, ordered=True)
            i += 1
            continue
        if stripped.startswith("- "):
            add_list_item(doc, stripped[2:].strip(), ordered=False)
            i += 1
            continue
        # Markdown hard-breaks.
        stripped = stripped.replace("  ", " ")
        add_paragraph(doc, stripped)
        i += 1

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("公司业务三个月 AI 改造行动计划｜内部讨论稿")
    cn_font(r, 8.5, False, GRAY)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
